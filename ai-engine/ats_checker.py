"""
ATS compatibility checks (font, alignment, format) for an uploaded resume.

The rules follow common resume guidelines:
  - Fonts: Arial, Calibri, Aptos or Times New Roman (other plain fonts tolerated)
  - Sizes: name 18-24 pt, section headings 12-14 pt, body 10-12 pt
  - Line spacing 1.0-1.15, margins 0.5-1 inch, left/justified body text
  - Simple single-column layout, standard sections, 1-2 pages
"""
import logging
import re
from collections import Counter

# pdfminer logs harmless font warnings for many browser-printed PDFs
logging.getLogger("pdfminer").setLevel(logging.ERROR)

def _font_key(name):
    """Compare font names ignoring case and spaces (PDFs store "TimesNewRoman")."""
    return name.lower().replace(" ", "")


RECOMMENDED_FONTS = {_font_key(f) for f in ("Arial", "Calibri", "Aptos", "Times New Roman")}
ACCEPTABLE_FONTS = RECOMMENDED_FONTS | {_font_key(f) for f in (
    "Helvetica", "Georgia", "Garamond", "Cambria", "Verdana", "Tahoma",
    "Book Antiqua", "Palatino Linotype", "Segoe UI", "Liberation Sans", "Carlito")}
# Bullet/icon fonts carry no real text, so they are left out of the font rules
SYMBOL_FONTS = {_font_key(f) for f in (
    "Symbol", "Wingdings", "Webdings", "Zapf Dingbats", "Segoe UI Emoji", "Segoe UI Symbol", "Unknown")}
SECTION_WORDS = ("summary", "objective", "profile", "skill", "experience", "education",
                 "project", "certification", "achievement", "internship", "language",
                 "interest", "hobbies", "declaration", "training", "award", "strength")
REQUIRED_SECTIONS = {
    "Summary/Objective": ("summary", "objective", "profile"),
    "Skills": ("skill",),
    "Education": ("education", "academic", "qualification"),
    "Experience/Projects": ("experience", "project", "internship", "employment"),
}
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
PHONE_RE = re.compile(r"\+?\d[\d\s()-]{8,}\d")
# Emoji / pictographic icons that ATS parsers often turn into garbage
ICON_RE = re.compile(r"[🌀-🫿☀-➿]")


def _result(score, issues, passes):
    score = max(0, min(100, round(score)))
    status = "pass" if score >= 80 else "warn" if score >= 50 else "fail"
    return {"score": score, "status": status, "issues": issues, "passes": passes}


def _is_section_heading(text):
    t = text.strip()
    return 0 < len(t) < 40 and (t.isupper() or t.istitle()) and any(w in t.lower() for w in SECTION_WORDS)


def _content_checks(text):
    """Format rules that only need the plain text (sections, contact info, length, icons)."""
    penalty, issues, passes = 0, [], []
    lower = text.lower()
    headings = [l.lower() for l in text.split("\n") if _is_section_heading(l)]
    heading_text = " ".join(headings) or lower

    missing = [name for name, words in REQUIRED_SECTIONS.items()
               if not any(w in heading_text for w in words)]
    if missing:
        penalty += 8 * len(missing)
        issues.append(f"Missing standard section(s): {', '.join(missing)}")
    else:
        passes.append("All standard sections present")

    has_email, has_phone = bool(EMAIL_RE.search(text)), bool(PHONE_RE.search(text))
    if has_email and has_phone:
        passes.append("Email and phone detected")
    else:
        penalty += 10
        issues.append("Contact info incomplete: " + ", ".join(
            n for n, ok in (("email", has_email), ("phone", has_phone)) if not ok) + " not found")

    words = len(text.split())
    if words < 250:
        penalty += 10
        issues.append(f"Resume is short ({words} words) - add more detail")
    elif words > 1100:
        penalty += 10
        issues.append(f"Resume is long ({words} words) - keep to 1-2 pages")
    else:
        passes.append(f"Good length ({words} words)")

    icons = len(ICON_RE.findall(text))
    if icons:
        penalty += min(10, 3 * icons)
        issues.append(f"{icons} icon/emoji character(s) - ATS may not read them")
    return penalty, issues, passes


# ---------------------------------------------------------------- DOCX ---

def _docx_checks(path, text):
    import docx
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = docx.Document(path)
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    default_rpr = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr")) if defaults is not None else None
    default_ppr = defaults.find(qn("w:pPrDefault") + "/" + qn("w:pPr")) if defaults is not None else None

    # Theme fonts (e.g. Calibri/Aptos) used when a style says "minorHAnsi" instead of a name
    theme_fonts = {}
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/theme"):
            blob = rel.target_part.blob.decode("utf-8", errors="ignore")
            for kind in ("major", "minor"):
                m = re.search(rf"<a:{kind}Font>\s*<a:latin typeface=\"([^\"]+)\"", blob)
                if m:
                    theme_fonts[kind] = m.group(1)

    def rpr_font(rpr):
        if rpr is None:
            return None
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            return None
        name = rfonts.get(qn("w:ascii")) or rfonts.get(qn("w:hAnsi"))
        if name:
            return name
        theme = rfonts.get(qn("w:asciiTheme")) or rfonts.get(qn("w:hAnsiTheme")) or ""
        return theme_fonts.get("major" if theme.startswith("major") else "minor") if theme else None

    def rpr_size(rpr):
        if rpr is None:
            return None
        sz = rpr.find(qn("w:sz"))
        return int(sz.get(qn("w:val"))) / 2 if sz is not None else None

    def style_chain(style):
        while style is not None:
            yield style
            style = style.base_style

    def resolve(run, para, getter, fallback):
        candidates = [run._r.rPr]
        if run.style is not None:
            candidates += [s.element.rPr for s in style_chain(run.style)]
        candidates += [s.element.rPr for s in style_chain(para.style)]
        candidates.append(default_rpr)
        for rpr in candidates:
            value = getter(rpr)
            if value:
                return value
        return fallback

    def para_alignment(para):
        for fmt in [para.paragraph_format] + [s.paragraph_format for s in style_chain(para.style)]:
            if fmt.alignment is not None:
                return fmt.alignment
        return WD_ALIGN_PARAGRAPH.LEFT

    def para_spacing(para):
        for fmt in [para.paragraph_format] + [s.paragraph_format for s in style_chain(para.style)]:
            if fmt.line_spacing is not None:
                return fmt.line_spacing
        spacing = default_ppr.find(qn("w:spacing")) if default_ppr is not None else None
        if spacing is not None and spacing.get(qn("w:line")) and spacing.get(qn("w:lineRule"), "auto") == "auto":
            return int(spacing.get(qn("w:line"))) / 240
        return 1.0

    # Gather per-paragraph font info
    paras = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        runs = [r for r in para.runs if r.text.strip()]
        fonts, sizes = Counter(), Counter()
        for r in runs:
            n = len(r.text.strip())
            fonts[resolve(r, para, rpr_font, "Times New Roman")] += n
            sizes[resolve(r, para, rpr_size, 10.0)] += n
        paras.append({"para": para, "text": t, "fonts": fonts, "sizes": sizes,
                      "size": max(sizes) if sizes else None})

    font_chars, size_chars = Counter(), Counter()
    for p in paras:
        font_chars.update({f: c for f, c in p["fonts"].items() if _font_key(f) not in SYMBOL_FONTS})
    # ---- Font check
    score, issues, passes = 100, [], []
    total = sum(font_chars.values()) or 1
    families = [f for f, _ in font_chars.most_common()]
    main_font = families[0] if families else "Unknown"
    unsafe = {f: c for f, c in font_chars.items() if _font_key(f) not in ACCEPTABLE_FONTS}
    if _font_key(main_font) in RECOMMENDED_FONTS:
        passes.append(f"Main font is {main_font}")
    elif _font_key(main_font) in ACCEPTABLE_FONTS:
        score -= 10
        issues.append(f"Main font {main_font} is OK, but Arial/Calibri/Aptos/Times New Roman is preferred")
    if unsafe:
        share = sum(unsafe.values()) / total
        score -= 15 + 35 * share
        issues.append(f"Non-ATS-friendly font(s): {', '.join(unsafe)} ({share:.0%} of text)")
    if len(families) > 2:
        score -= 10
        issues.append(f"{len(families)} different fonts used ({', '.join(families[:4])}) - stick to 1-2")

    name = paras[0] if paras else None
    headings = [p for p in paras[1:] if _is_section_heading(p["text"]) and p["size"]]
    body = [p for p in paras[1:] if p not in headings]
    for p in body:
        size_chars.update(p["sizes"])
    body_size = size_chars.most_common(1)[0][0] if size_chars else None

    if body_size is not None:
        if 10 <= body_size <= 12:
            passes.append(f"Body text {body_size:g} pt (10-12 pt)")
        else:
            score -= 20
            issues.append(f"Body text is {body_size:g} pt - use 10-12 pt (11 pt ideal)")
        off = sum(c for s, c in size_chars.items() if s < 9 or s > 14) / (sum(size_chars.values()) or 1)
        if off > 0.1:
            score -= 10
            issues.append(f"{off:.0%} of body text is below 9 pt or above 14 pt")
    if name and name["size"]:
        if 18 <= name["size"] <= 24:
            passes.append(f"Name {name['size']:g} pt (18-24 pt)")
        else:
            score -= 10
            issues.append(f"Name is {name['size']:g} pt - use 18-24 pt (22 pt bold ideal)")
    if headings:
        bad = sorted({p["size"] for p in headings if not 12 <= p["size"] <= 14})
        if bad:
            score -= 10
            issues.append(f"Section headings at {', '.join(f'{s:g}' for s in bad)} pt - use 12-14 pt (13 pt bold)")
        else:
            passes.append("Section headings 12-14 pt")
    font_result = _result(score, issues, passes)
    font_result["detected"] = {"main_font": main_font, "body_size": body_size}

    # ---- Alignment check
    score, issues, passes = 100, [], []
    # The top block (name, title, contact) may be centred; everything after must be left/justified
    first_section = next((i for i, p in enumerate(paras) if _is_section_heading(p["text"])), len(paras))
    body_paras = [p for p in paras[first_section:] if not _is_section_heading(p["text"])]
    aligns = Counter(para_alignment(p["para"]) for p in body_paras)
    ok_aligns = {WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY}
    bad_align = sum(c for a, c in aligns.items() if a not in ok_aligns)
    if body_paras:
        if bad_align:
            share = bad_align / len(body_paras)
            score -= 10 + 30 * share
            issues.append(f"{bad_align} body paragraph(s) centred/right-aligned - keep body text left or justified")
        else:
            passes.append("Body text left/justified")
        if aligns[WD_ALIGN_PARAGRAPH.LEFT] and aligns[WD_ALIGN_PARAGRAPH.JUSTIFY]:
            minority = min(aligns[WD_ALIGN_PARAGRAPH.LEFT], aligns[WD_ALIGN_PARAGRAPH.JUSTIFY])
            if minority / len(body_paras) > 0.15:
                score -= 5
                issues.append("Mix of left and justified paragraphs - pick one")

    head_aligns = {para_alignment(p["para"]) for p in paras[:first_section]}
    if len(head_aligns) > 1:
        score -= 5
        issues.append("Name/title/contact lines are not aligned the same way")

    spacings = Counter()
    for p in paras:
        ls = para_spacing(p["para"])
        spacings[round(ls, 2) if isinstance(ls, float) else "exact"] += 1
    main_spacing = spacings.most_common(1)[0][0] if spacings else 1.0
    if main_spacing == "exact":
        passes.append("Fixed line spacing used")
    elif 1.0 <= main_spacing <= 1.2:
        passes.append(f"Line spacing {main_spacing:g} (1.0-1.15)")
    elif main_spacing <= 1.5:
        score -= 8
        issues.append(f"Line spacing {main_spacing:g} - 1.0-1.15 is recommended")
    else:
        score -= 15
        issues.append(f"Line spacing {main_spacing:g} is too loose - use 1.0-1.15")

    for i, sec in enumerate(doc.sections):
        margins = {"left": sec.left_margin, "right": sec.right_margin,
                   "top": sec.top_margin, "bottom": sec.bottom_margin}
        bad = {k: v.inches for k, v in margins.items() if v is not None and not 0.49 <= v.inches <= 1.01}
        if bad:
            score -= 5 * len(bad)
            issues.append("Margins outside 0.5-1 inch: " + ", ".join(f"{k} {v:.2f}\"" for k, v in bad.items()))
        elif i == 0:
            passes.append("Margins within 0.5-1 inch")

    indents = {round(p["para"].paragraph_format.left_indent.inches, 2)
               for p in body_paras if p["para"].paragraph_format.left_indent}
    if len(indents) > 3:
        score -= 8
        issues.append(f"{len(indents)} different indent levels - keep indentation consistent")
    manual = sum(1 for p in paras if re.search(r"\S {4,}\S|\t\S.*\t", p["para"].text))
    if manual:
        score -= min(15, 3 * manual)
        issues.append(f"{manual} line(s) aligned with spaces/tabs - ATS may merge or split the text")
    align_result = _result(score, issues, passes)

    # ---- Format check
    penalty, issues, passes = _content_checks(text)
    score = 100 - penalty
    body_xml = doc.element.body.xml
    if doc.tables:
        score -= 15
        issues.append(f"{len(doc.tables)} table(s) - ATS often reads tables out of order")
    if "txbxContent" in body_xml:
        score -= 20
        issues.append("Text boxes found - ATS usually skips text inside them")
    cols = re.search(r'<w:cols [^>]*w:num="(\d+)"', body_xml)
    if cols and int(cols.group(1)) > 1:
        score -= 20
        issues.append("Multi-column layout - use a single column")
    images = body_xml.count("<w:drawing") + body_xml.count("<w:pict")
    if images:
        score -= 5
        issues.append(f"{images} image/graphic(s) - ATS ignores them")
    header_text = " ".join(p.text for s in doc.sections for part in (s.header, s.footer)
                           for p in part.paragraphs if not part.is_linked_to_previous).strip()
    if EMAIL_RE.search(header_text) or PHONE_RE.search(header_text):
        score -= 10
        issues.append("Contact details are in the page header/footer - ATS may miss them")
    if not (doc.tables or cols or "txbxContent" in body_xml):
        passes.append("Simple single-column layout")
    passes.append("DOCX file (ATS-friendly)")
    return font_result, align_result, _result(score, issues, passes)


# ----------------------------------------------------------------- PDF ---

def _pdf_checks(path, text):
    from PyPDF2 import PdfReader
    from pdfminer.high_level import extract_pages
    from pdfminer.layout import LTChar, LTTextLine

    reader = PdfReader(path)
    lines = []  # one entry per text line, with exact positions from pdfminer
    page_w = 612

    def walk(obj):
        if isinstance(obj, LTTextLine):
            yield obj
        elif hasattr(obj, "__iter__"):
            for child in obj:
                yield from walk(child)

    for page_no, layout in enumerate(extract_pages(path)):
        if page_no == 0:
            page_w = layout.width
        for line in walk(layout):
            chars = [c for c in line if isinstance(c, LTChar) and c.get_text().strip()]
            t = line.get_text().strip()
            if not chars or not t:
                continue
            fonts = Counter(re.sub(r"^[A-Z]{6}\+", "", c.fontname) for c in chars)
            sizes = Counter(round(c.size, 1) for c in chars)
            lines.append({"page": page_no, "x0": line.x0, "x1": line.x1, "y": line.y0,
                          "fonts": fonts, "size": sizes.most_common(1)[0][0], "text": t})

    def family(font):
        base = re.split(r"[-,]", font)[0]
        base = re.sub(r"(MT|PS|Std|Bold|Italic|Regular)+$", "", base)
        return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", base).strip() or "Unknown"

    # ---- Font check
    score, issues, passes = 100, [], []
    font_chars, size_chars = Counter(), Counter()
    for ln in lines:
        for font, count in ln["fonts"].items():
            if _font_key(family(font)) not in SYMBOL_FONTS:
                font_chars[family(font)] += count
        size_chars[ln["size"]] += len(ln["text"])
    total = sum(font_chars.values()) or 1
    families = [f for f, _ in font_chars.most_common()]
    main_font = families[0] if families else "Unknown"
    unsafe = {f: c for f, c in font_chars.items() if _font_key(f) not in ACCEPTABLE_FONTS}
    if _font_key(main_font) in RECOMMENDED_FONTS:
        passes.append(f"Main font is {main_font}")
    elif _font_key(main_font) in ACCEPTABLE_FONTS:
        score -= 10
        issues.append(f"Main font {main_font} is OK, but Arial/Calibri/Aptos/Times New Roman is preferred")
    if unsafe:
        share = sum(unsafe.values()) / total
        score -= 15 + 35 * share
        issues.append(f"Non-ATS-friendly font(s): {', '.join(list(unsafe)[:3])} ({share:.0%} of text)")
    if len(families) > 2:
        score -= 10
        issues.append(f"{len(families)} different fonts used ({', '.join(families[:4])}) - stick to 1-2")

    body_size = size_chars.most_common(1)[0][0] if size_chars else None
    if body_size:
        if 10 <= body_size <= 12:
            passes.append(f"Body text {body_size:g} pt (10-12 pt)")
        else:
            score -= 20
            issues.append(f"Body text is {body_size:g} pt - use 10-12 pt (11 pt ideal)")
    first_page = [ln for ln in lines if ln["page"] == 0]
    if first_page:
        name_size = max(ln["size"] for ln in first_page)
        if 18 <= name_size <= 24:
            passes.append(f"Name {name_size:g} pt (18-24 pt)")
        else:
            score -= 10
            issues.append(f"Largest text (name) is {name_size:g} pt - use 18-24 pt (22 pt bold ideal)")
    heading_sizes = sorted({ln["size"] for ln in lines if _is_section_heading(ln["text"])})
    bad = [s for s in heading_sizes if not 12 <= s <= 14]
    if heading_sizes and bad:
        score -= 10
        issues.append(f"Section headings at {', '.join(f'{s:g}' for s in bad)} pt - use 12-14 pt (13 pt bold)")
    elif heading_sizes:
        passes.append("Section headings 12-14 pt")
    font_result = _result(score, issues, passes)
    font_result["detected"] = {"main_font": main_font, "body_size": body_size}

    # ---- Alignment check (from text positions)
    score, issues, passes = 100, [], []
    # Body lines and section headings; skips the name block and tiny browser header/footer text
    first_section = next((i for i, ln in enumerate(lines) if _is_section_heading(ln["text"])), 0)
    body = [ln for ln in lines[first_section:] if body_size and ln["size"] >= body_size - 0.6]
    if body:
        left_edge = Counter(round(ln["x0"]) for ln in body).most_common(1)[0][0]
        left_margin = left_edge / 72
        if 0.49 <= left_margin <= 1.01:
            passes.append(f"Left margin {left_margin:.2f}\" (0.5-1 inch)")
        else:
            score -= 10
            issues.append(f"Left margin is about {left_margin:.2f}\" - use 0.5-1 inch")

        rows = {}
        for ln in body:
            rows.setdefault((ln["page"], round(ln["y"] / 3)), []).append(ln)
        centred, right, split_rows, line_starts = 0, 0, 0, Counter()
        for row in rows.values():
            row.sort(key=lambda ln: ln["x0"])
            start, end = row[0]["x0"], row[-1]["x1"]
            if len(row) > 1:
                split_rows += 1  # text pieces side by side: columns or tab-aligned dates
            if start > left_edge + 0.15 * page_w and abs((start + end) / 2 - page_w / 2) < page_w * 0.06:
                centred += 1
            elif start > page_w * 0.5:
                right += 1
            else:
                line_starts[round(start / 6)] += 1

        # Left-aligned lines should begin at a few consistent x positions (margin + bullet indents)
        common = sum(c for _, c in line_starts.most_common(3))
        ragged = 1 - common / (sum(line_starts.values()) or 1)
        if ragged > 0.2:
            score -= 10 + 30 * ragged
            issues.append(f"{ragged:.0%} of lines start at irregular positions - keep a consistent left edge")
        else:
            passes.append("Body lines start at a consistent left edge")
        if centred:
            score -= min(15, 3 * centred)
            issues.append(f"{centred} centred line(s) after the header - keep headings and text left-aligned")
        if right:
            score -= min(10, 2 * right)
            issues.append(f"{right} right-aligned line(s) - ATS may attach them to the wrong section")
        if split_rows > 3:
            score -= 10
            issues.append(f"{split_rows} lines have text spread across the page (columns/tabs) - may be read out of order")

        # Line spacing: median baseline gap / font size (Word "single" is about 1.15 x size)
        ys = sorted({(ln["page"], round(ln["y"], 1)) for ln in body if abs(ln["size"] - body_size) < 0.6},
                    key=lambda k: (k[0], -k[1]))
        gaps = [a[1] - b[1] for a, b in zip(ys, ys[1:]) if a[0] == b[0] and 0 < a[1] - b[1] < body_size * 3]
        if gaps:
            ratio = sorted(gaps)[len(gaps) // 2] / body_size / 1.15
            if ratio <= 1.25:
                passes.append(f"Line spacing about {max(ratio, 1.0):.2f}")
            else:
                score -= 8 if ratio <= 1.6 else 15
                issues.append(f"Line spacing about {ratio:.2f} - 1.0-1.15 is recommended")
    else:
        score -= 20
        issues.append("Could not read text positions from this PDF")
    align_result = _result(score, issues, passes)

    # ---- Format check
    penalty, issues, passes = _content_checks(text)
    score = 100 - penalty
    pages = len(reader.pages)
    if pages > 2:
        score -= 15
        issues.append(f"{pages} pages - keep to 1 page (fresher) or 2 pages")
    else:
        passes.append(f"{pages} page(s)")
    images = 0
    for page in reader.pages:
        xobjects = (page.get("/Resources") or {}).get("/XObject") or {}
        try:
            images += sum(1 for o in xobjects.values() if o.get_object().get("/Subtype") == "/Image")
        except Exception:
            pass
    if images:
        score -= 5
        issues.append(f"{images} image(s) - ATS ignores them")
    if len(text.strip()) < 200:
        score -= 30
        issues.append("Very little selectable text - the PDF may be scanned/image-based")
    passes.append("PDF with selectable text")
    return font_result, align_result, _result(score, issues, passes)


def analyze_ats(path, text):
    """Return {'font': {...}, 'alignment': {...}, 'format': {...}} or None if unsupported."""
    try:
        if path.lower().endswith(".docx"):
            font, align, fmt = _docx_checks(path, text)
        elif path.lower().endswith(".pdf"):
            font, align, fmt = _pdf_checks(path, text)
        else:
            return None
    except Exception as exc:
        return {"error": f"ATS check failed: {exc}"}
    return {"font": font, "alignment": align, "format": fmt}
